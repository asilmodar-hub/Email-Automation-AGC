"""Read the contact list, and decide who is written to in a round.

The workbook is only ever read. Headers are matched by meaning, in German or
English, so a list AGC already keeps works without being reshaped. That covers
the contact template that ships with the tool and AGC's own Deal Origination
Cockpit (sheet "Intermediäre": Firma, Kategorie, Ansprechpartner, E-Mail, Ort,
Region, Zuständigkeit, Prio. and the rest; see references/08-agc-contact-list.md).
Every sheet with an email column and at least one usable address is read, and
every column is kept: what has no field of its own travels with the contact as
background for Claude.
"""
import datetime as dt
import json
import os
import re
import unicodedata
import warnings

import openpyxl

from paths import long_path

HERE = os.path.dirname(os.path.abspath(__file__))
DIRECTORY_FILE = os.path.join(os.path.dirname(HERE), "assets", "agc_directory.json")

# field -> header spellings that mean it (compared lower-case, punctuation removed)
FIELDS = {
    "first": ["first name", "firstname", "vorname", "given name"],
    "last": ["last name", "lastname", "surname", "nachname", "familienname", "name"],
    "person": ["ansprechpartner", "ansprechpartnerin", "ansprechperson", "ansprechpersonen", "ansprechperson en",
               "kontaktperson", "contact person", "contact", "person", "full name"],
    "title": ["title", "titel", "akad titel", "academic title"],
    "gender": ["salutation", "anrede", "gender", "geschlecht", "mr ms", "herr frau"],
    "company": ["company", "firma", "unternehmen", "organisation", "organization", "kanzlei", "firm", "bank"],
    "role": ["role", "position", "funktion", "job title", "rolle"],
    "department": ["bereich", "abteilung", "department"],
    "type": ["type", "contact type", "category", "kategorie", "typ", "kontakttyp", "art", "kontaktart"],
    "email": ["email", "e mail", "mail", "email address", "e mail adresse", "emailadresse", "mailadresse"],
    "language": ["language", "sprache", "lang"],
    "form": ["form", "sie du", "du sie", "address form", "anredeform", "formality"],
    "owner": ["owner", "agc owner", "agc contact", "betreuer", "verantwortlich", "agc ansprechpartner",
              "relationship owner", "zuständigkeit", "zustaendigkeit", "zuständig", "zustaendig", "zuständig agc"],
    "relationship": ["relationship", "beziehung", "how we know them", "bekannt durch", "kontext", "context"],
    "sector": ["sector focus", "sector", "branche", "branchenfokus", "focus", "fokus", "schwerpunkt",
               "spezialexpertise"],
    "notes": ["notes", "note", "notizen", "notiz", "bemerkung", "bemerkungen", "anmerkung", "anmerkungen",
              "comment", "comments", "kommentar"],
    "todos": ["sonstige todos", "todos", "to dos", "aufgaben"],
    "cadence": ["cadence", "rhythm", "rhythmus", "frequency", "frequenz", "turnus"],
    "dnc": ["do not contact", "dnc", "nicht kontaktieren", "opt out", "optout", "kein kontakt"],
    "greeting": ["greeting", "greeting override", "begruessung", "begrüßung", "grussformel"],
    "region": ["region", "geografie", "geography", "land", "country"],
    "city": ["ort", "stadt", "city", "standort", "nächster standort", "town"],
    "postcode": ["plz", "postleitzahl", "zip", "postcode", "postal code"],
    "address": ["adresse", "address", "anschrift", "straße", "strasse"],
    "phone": ["tel", "telefon", "phone", "telephone", "mobil", "mobile"],
    "website": ["internet", "website", "web", "homepage", "url"],
    "priority": ["prio", "priorität", "prioritaet", "priority", "ranking"],
    "relevance": ["relevanz für at deal flow", "relevanz", "relevance"],
    "first_contact_date": ["datum erstansprache", "date first contact", "first contacted"],
}
# Academic titles and the like, as they are written around a name. Those that
# go in front of the name in a greeting are kept; letters after it are dropped.
PRE_TITLES = {"univ.-prof.": "Univ.-Prof.", "univ.prof.": "Univ.-Prof.", "prof.": "Prof.", "prof": "Prof.",
              "dr.": "Dr.", "dr": "Dr.", "mag.": "Mag.", "mag": "Mag.", "mmag.": "MMag.", "mmag": "MMag.",
              "dkfm.": "Dkfm.", "dipl.-ing.": "DI", "dipl.-kfm.": "Dipl.-Kfm.", "dipl.-kffr.": "Dipl.-Kffr.",
              "dipl.-vw.": "Dipl.-Vw.", "ing.": "Ing.", "lic.": "lic.", "dipl.": "Dipl.", "mag.(fh)": "Mag. (FH)"}
POST_TITLES = {"mba", "msc", "bsc", "ll.m.", "llm", "cfa", "cpa", "phd", "ph.d.", "b.a.", "m.a.", "(fh)", "jun.",
               "sen.", "jr.", "sr."}
PARTICLES = {"von", "van", "de", "der", "den", "zu", "zum", "dos", "da", "di", "del", "la", "le", "ten", "ter"}
GENERIC_MAIL = {"gmail.com", "gmx.at", "gmx.de", "gmx.net", "yahoo.com", "yahoo.de", "hotmail.com", "outlook.com",
                "icloud.com", "me.com", "web.de", "aon.at", "chello.at", "a1.net", "live.com", "t-online.de",
                "googlemail.com", "proton.me", "protonmail.com"}
REQUIRED = ["email", "name"]
# what the tool works out itself when the list does not say it
WORKED_OUT = ("gender", "language", "form")

# What kind of contact someone is, for the chart on the Contacts tab. Taken from
# the Type column when there is one; otherwise read off the firm's name and the
# role. First match wins, so the narrow words come before the broad ones.
TYPE_RULES = [
    ("Law firm", ["rechtsanw", "anwalt", "anwält", "kanzlei", "legal", "law", "attorney", "notar", "solicitor"]),
    ("Tax / audit", ["steuerberat", "wirtschaftsprüf", "wirtschaftspruef", "tax", "audit", "treuhand",
                     "accounting", "accountant", "wp/stb"]),
    ("M&A adviser", ["m&a", "corporate finance", "mergers", "transaction", "investment bank", "investmentbank"]),
    ("Venture capital", ["venture", " vc", "vc ", "seed fund"]),
    ("Private equity", ["private equity", "beteiligung", "buyout", "buy-out", " pe ", "family office",
                        "finanzinvestor", "fund manager", "fof"]),
    ("Bank", ["bank", "sparkasse", "raiffeisen", "volksbank", "kreditinstitut"]),
    ("Consultant", ["consult", "berat", "advisory", "adviser", "advisor", "interim"]),
]
OTHER_TYPE = "Other"

CADENCE_DAYS = {"quarterly": 80, "half_yearly": 170, "yearly": 350}
RECENT_GUARD_DAYS = 30

EMAIL_RE = re.compile(r"^[^@\s,;<>]+@[^@\s,;<>]+\.[^@\s,;<>]{2,}$")
FIND_EMAIL = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}")
CODE_RE = re.compile(r"^[A-ZÄÖÜ]{2,3}$")
PLACEHOLDER_ROW = re.compile(r"^\s*(\[.*\]|tbd|n\.?\s?a\.?|#n/a|-+|x+)\s*$", re.IGNORECASE)

# Problems that a look online can solve: the person writing still wants this
# firm, so Claude finds the missing piece and the reviewer decides on it.
RESEARCHABLE = {"no email address": "email", "email address looks wrong": "email", "no name": "person"}


def _type_of(text):
    low = " %s " % str(text or "").lower()
    for name, words in TYPE_RULES:
        if any(w in low for w in words):
            return name
    return ""


def contact_type(given, company, role):
    """(type, guessed). A value in the Type column always wins: a known kind is
    put under its usual name, anything else is kept as the list has it."""
    given = str(given or "").strip()
    if given:
        return (_type_of(given) or given[:1].upper() + given[1:]), False
    return (_type_of(company) or _type_of(role) or OTHER_TYPE), True


_DIRECTORY = None


def directory():
    """assets/agc_directory.json: the team's short forms, countries and regions."""
    global _DIRECTORY
    if _DIRECTORY is None:
        try:
            with open(DIRECTORY_FILE, encoding="utf-8") as fh:
                _DIRECTORY = json.load(fh)
        except (OSError, ValueError):
            _DIRECTORY = {}
    return _DIRECTORY


def _norm(text):
    return re.sub(r"[^a-z0-9äöüß]+", " ", str(text or "").lower()).strip()


def _ascii(text):
    """'Kröhan' -> 'kroehan', for comparing a name with an email address."""
    t = str(text or "").lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        t = t.replace(a, b)
    t = unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z]", "", t)


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\xa0", " ").strip()
    return "" if text in ("|", "#N/A", "#NV") else text


def _gender(text):
    t = _norm(text)
    if t in ("herr", "mr", "m", "male", "mann", "männlich", "maennlich", "hr"):
        return "m"
    if t in ("frau", "ms", "mrs", "f", "w", "female", "weiblich", "fr"):
        return "f"
    return ""


def _language(text):
    t = _norm(text)
    if t.startswith(("en", "eng")):
        return "en"
    if t.startswith(("de", "ger", "deu")):
        return "de"
    return ""


def _form(text):
    t = _norm(text)
    if t in ("du", "informal", "first name", "vorname", "per du", "informell"):
        return "du"
    if t in ("sie", "formal", "per sie", "formell"):
        return "sie"
    return ""


def _cadence(text):
    t = _norm(text)
    if not t:
        return ""
    if t.startswith(("q", "viertel")) or "quart" in t:
        return "quarterly"
    if t.startswith(("half", "halb", "semi")) or "6" in t:
        return "half_yearly"
    if t.startswith(("year", "jähr", "jaehr", "annual")) or "12" in t:
        return "yearly"
    if t.startswith(("never", "nie", "none", "kein")):
        return "never"
    return ""


def _yes(text):
    return _norm(text) in ("x", "yes", "ja", "y", "j", "true", "1", "wahr")


def _priority(text):
    t = str(text or "").strip().upper().rstrip(".")
    return t if t in directory().get("priorities", {}) else ""


def _relevance(text):
    t = _norm(text)
    return {"hoch": "high", "high": "high", "mittel": "medium", "medium": "medium", "niedrig": "low",
            "low": "low"}.get(t, "")


# ---- where someone is -------------------------------------------------------------

def _has_word(haystack, word):
    return re.search(r"(?<![a-zäöüß])%s(?![a-zäöüß])" % re.escape(word), haystack) is not None


def _country_of(text):
    """Country codes a free-text region or city names, e.g. 'Hessen' -> ['DE']."""
    low = " ".join(str(text or "").lower().replace("\n", " ").split())
    if not low:
        return []
    d = directory()
    out = []
    for name, codes in d.get("multi", {}).items():
        if _has_word(low, name.lower()):
            out.extend(c for c in codes if c not in out)
    for code, words in d.get("countries", {}).items():
        if code not in out and any(_has_word(low, w) for w in words):
            out.append(code)
    return out


def _state_of(country, text):
    low = " ".join(str(text or "").lower().split())
    for state, words in directory().get("states", {}).get(country, {}).items():
        if any(_has_word(low, w) for w in words):
            return state
    return ""


def place(region, city, postcode, email):
    """(country, state, how): where a contact is, from the Region column, the
    city, the postcode and in the end the address's country ending."""
    for source, text in (("region", region), ("city", city)):
        found = _country_of(text)
        if len(found) == 1:
            return found[0], _state_of(found[0], text) or _state_of(found[0], city), source
        if len(found) > 1:
            return "MULTI", "", source
    code = str(postcode or "").strip()
    domain = email.rsplit("@", 1)[-1] if "@" in email else ""
    tld = domain.rsplit(".", 1)[-1].lower() if domain else ""
    by_tld = directory().get("tld", {}).get(tld, "")
    if re.fullmatch(r"\d{5}", code):
        return "DE", "", "postcode"
    if re.fullmatch(r"\d{4}", code) and by_tld in ("AT", "CH"):
        return by_tld, "", "postcode"
    if by_tld:
        return by_tld, "", "email"
    return "", "", ""


# ---- names --------------------------------------------------------------------

def _strip_titles(text):
    """'Dr. Jürgen Kuttruff' -> ('Jürgen Kuttruff', 'Dr.')"""
    kept, titles = [], []
    for w in str(text or "").replace(",", " ").split():
        low = w.lower()
        if low in PRE_TITLES or w == "DI":
            titles.append(PRE_TITLES.get(low, w))
        elif low not in POST_TITLES:
            kept.append(w)
    return " ".join(kept), " ".join(titles)


def _proper(word):
    return word[:1].upper() + word[1:].lower() if word.isupper() and len(word) > 2 else word


def _one_person(text):
    """'Dr. Robert Kremlicka (Vice President)' -> {first, last, title, role}"""
    role = " ".join(re.findall(r"\(([^)]*)\)", text)).strip()
    text = re.sub(r"\([^)]*\)", " ", text)
    text = re.sub(r"\s+", " ", text).strip(" ,.-")
    gender = ""
    m = re.match(r"(?i)^(herr|hr\.?|mr\.?|frau|fr\.?|ms\.?|mrs\.?)\s+", text)
    if m:
        gender = "m" if m.group(1).lower().rstrip(".") in ("herr", "hr", "mr") else "f"
        text = text[m.end():]
    name, title = _strip_titles(text)
    words = [_proper(w) for w in name.split() if w]
    if not words:
        return None
    if len(words) == 1:
        return {"first": "", "last": words[0], "title": title, "role": role, "gender": gender}
    last = [words[-1]]
    i = len(words) - 2
    while i > 0 and words[i].lower() in PARTICLES:
        last.insert(0, words[i])
        i -= 1
    return {"first": " ".join(words[:i + 1]), "last": " ".join(last), "title": title, "role": role,
            "gender": gender}


def people(text, email=""):
    """The people a cell names. 'Mäkki, Andreas Huber' is two people;
    'Barbist, Johannes' one, written surname first, unless the address
    says otherwise."""
    text = str(text or "").replace("\r", "\n").strip()
    if not text or PLACEHOLDER_ROW.match(text):
        return []
    local = _ascii(email.split("@")[0]) if "@" in email else ""
    commas = text.count(",")
    other_seps = re.search(r"[;/&\n]| und | and ", text)
    if commas == 1 and not other_seps:
        left, right = [p.strip() for p in text.split(",")]
        l_name, l_title = _strip_titles(left)
        r_name, r_title = _strip_titles(right)
        if len(l_name.split()) == 1 and len(r_name.split()) == 1 and r_name:
            if not local or _ascii(l_name) in local:
                one = _one_person(" ".join(filter(None, [l_title or r_title, r_name, l_name])))
                return [one] if one else []
    parts = re.split(r"[,;/&\n]|\bund\b|\band\b", text)
    return [p for p in (_one_person(x) for x in parts) if p and (p["last"] or p["first"])]


def _initials_of(p):
    return (_ascii(p.get("first", ""))[:1] + _ascii(p.get("last", ""))[:1])


def pick(found, email):
    """The person an address belongs to, or the first one named."""
    if not found:
        return None, []
    local = _ascii(email.split("@")[0]) if "@" in email else ""
    best, score = found[0], -1
    for p in found:
        s = 0
        last, first = _ascii(p["last"]), _ascii(p["first"])
        if last and last in local:
            s += 3
        if first and first in local:
            s += 1
        if local and len(local) <= 3 and _initials_of(p) and local.startswith(_initials_of(p)):
            s += 1
        if s > score:
            best, score = p, s
    return best, [q for q in found if q is not best]


def _first_from_address(person, email):
    """'roman.pongracz@…' gives a first name the list does not."""
    local = email.split("@")[0].lower() if "@" in email else ""
    parts = [x for x in re.split(r"[._-]", local) if x]
    if len(parts) == 2 and _ascii(person["last"]) == _ascii(parts[1]) and len(parts[0]) > 2:
        return parts[0].capitalize()
    return ""


def emails_in(text):
    return [m.group(0).strip(".").lower() for m in FIND_EMAIL.finditer(str(text or ""))]


# ---- the workbook ---------------------------------------------------------------

def _header_key(value):
    # "Salutation (Herr/Frau)" is the Salutation column: hints in brackets are dropped
    return _norm(re.sub(r"\(.*?\)|\[.*?\]", " ", str(value or "").replace("\n", " ")))


def _find_header(rows):
    """The first row among the top ten that looks like the header row."""
    best = (0, None, {})
    for r, row in enumerate(rows[:10], start=1):
        mapping = {}
        for c, value in enumerate(row):
            key = _header_key(value)
            if not key:
                continue
            for field, names in FIELDS.items():
                if field not in mapping and key in names:
                    mapping[field] = c
                    break
        if len(mapping) > best[0]:
            best = (len(mapping), r, mapping)
    return best


def _owner_matrix(rows, header_row, mapping):
    """A row of short forms under the header (KLA, DJE, … with an X below for
    whoever looks after the contact) instead of one owner per row. Returns
    {column: code}, or {} when the row under the header is an ordinary one."""
    if header_row >= len(rows):
        return {}
    header, below = rows[header_row - 1], rows[header_row]
    owner_col = mapping.get("owner")
    codes = {}
    for c, v in enumerate(below):
        text = _cell(v)
        if not CODE_RE.match(text):
            continue
        if c == owner_col or c >= len(header) or not _cell(header[c]):
            codes[c] = text
    email_col = mapping.get("email")
    if len(codes) < 2 or (email_col is not None and email_col < len(below) and _cell(below[email_col])):
        return {}
    return codes


def _comments(ws):
    """Cell comments by row. Threaded comments carry Excel's notice first;
    only what the person wrote is kept."""
    out = {}
    try:
        cells = [c for row in ws.iter_rows() for c in row if c.comment]
    except Exception:
        return out
    for c in cells:
        text = c.comment.text or ""
        if "Kommentar:" in text or "Comment:" in text:
            text = re.split(r"Kommentar:|Comment:", text, maxsplit=1)[1]
        # a note starts with its author's name ("Christoph Ganahl:"); the rest is what they wrote
        text = re.sub(r"^[^\n:]{3,40}:\s*\n", "", text.strip() + "\n", count=1)
        text = " ".join(text.split())
        if text:
            out.setdefault(c.row, []).append(text)
    return out


def _sheet_contacts(ws, rows, header_row, mapping, matrix):
    header = rows[header_row - 1]
    first_data = header_row + (1 if matrix else 0)
    matched = set(mapping.values()) | set(matrix)
    names, extra_cols = {}, []
    for c, v in enumerate(header):
        name = " ".join(_cell(v).split())
        if c in matched or not name:
            continue
        if name in names.values():
            name = "%s (%s)" % (name, openpyxl.utils.get_column_letter(c + 1))
        names[c] = name
        extra_cols.append(c)
    comments = _comments(ws)
    out = []
    for r, row in enumerate(rows[first_data:], start=first_data + 1):
        get = lambda f: _cell(row[mapping[f]]) if f in mapping and mapping[f] < len(row) else ""
        if not any(_cell(v) for v in row):
            continue
        company, person_cell, email_cell = get("company"), get("person"), get("email")
        if company and PLACEHOLDER_ROW.match(company) and not email_cell and not person_cell:
            continue
        addresses = emails_in(email_cell)
        email = addresses[0] if addresses else ("" if not email_cell else email_cell.strip().lower())
        if not addresses and email_cell and not EMAIL_RE.match(email_cell.strip()):
            # a name, a city or "n.a." in the address column: keep it as a note, not as an address
            email = ""
        # the person: own columns first, then the "Ansprechpartner" cell, then a name in the address cell
        first, last, title, role_hint, others = get("first"), get("last"), get("title"), "", []
        gender_hint = ""
        if not first and not last:
            found = people(person_cell, email)
            if not found and email_cell and addresses:
                found = people(re.sub(FIND_EMAIL, " ", email_cell).replace("<", " ").replace(">", " "), email)
            chosen, rest = pick(found, email)
            if chosen:
                first, last, title, role_hint = chosen["first"], chosen["last"], title or chosen["title"], chosen["role"]
                gender_hint = chosen.get("gender", "")
                others = [" ".join(filter(None, [p["title"], p["first"], p["last"]])) for p in rest]
        elif not first and " " in last:
            # one "Name" column holding "First Last"
            first, last = last.rsplit(" ", 1)
        assumed = []
        if last and not first and email:
            guess = _first_from_address({"last": last}, email)
            if guess:
                first = guess
                assumed.append("first name taken from the email address")
        owner = get("owner").upper() if CODE_RE.match(get("owner").upper()) else get("owner")
        if matrix:
            owner = ", ".join(code for c, code in matrix.items() if c < len(row) and _cell(row[c]))
        country, state, where_from = place(get("region"), get("city"), get("postcode"), email)
        extra = {}
        for c in extra_cols:
            if c < len(row) and _cell(row[c]):
                extra[names[c]] = _cell(row[c])[:300]
        if len(addresses) > 1:
            extra["Further addresses"] = ", ".join(addresses[1:])
        if email_cell and not addresses:
            extra["Address column says"] = email_cell[:120]
        contact = {
            "row": r, "sheet": ws.title, "id": email or "%s:%d" % (_norm(ws.title).replace(" ", "-") or "row", r),
            "first": first, "last": last, "title": title,
            "gender": _gender(get("gender")) or gender_hint, "company": company, "role": get("role") or role_hint,
            "department": get("department"), "email": email, "language": _language(get("language")),
            "form": _form(get("form")), "owner": owner, "relationship": get("relationship"),
            "sector": get("sector"), "notes": get("notes"), "todos": get("todos"),
            "cadence": _cadence(get("cadence")), "dnc": _yes(get("dnc")), "greeting": get("greeting"),
            "region": get("region"), "city": get("city"), "postcode": get("postcode"), "address": get("address"),
            "phone": get("phone"), "website": get("website"), "country": country, "state": state,
            "priority": _priority(get("priority")), "relevance": _relevance(get("relevance")),
            "first_contact_date": get("first_contact_date") if re.match(r"\d{4}-\d\d-\d\d", get("first_contact_date")) else "",
            "others": others, "comments": comments.get(r, []), "extra": extra,
            "problems": [], "assumed": assumed,
        }
        contact["type"], contact["type_guessed"] = contact_type(get("type"), company, contact["role"])
        out.append(contact)
    return out


def read_workbook(path):
    """Return {ok, problem, sheet, sheets, contacts, missing_columns, worked_out, unmatched_headers, owners}."""
    out = {"ok": False, "problem": "", "sheet": "", "sheets": [], "contacts": [], "missing_columns": [],
           "worked_out": [], "unmatched_headers": [], "owners": {}, "path": path}
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")   # Excel features openpyxl does not know (data validation extensions)
            wb = openpyxl.load_workbook(long_path(path), data_only=True)
    except FileNotFoundError:
        out["problem"] = "The contact list was not found at this path."
        return out
    except PermissionError:
        out["problem"] = "The contact list could not be opened. Windows refused access to it."
        return out
    except Exception as exc:
        out["problem"] = "This file could not be read as an Excel workbook (%s)." % type(exc).__name__
        return out

    found = []
    for ws in wb.worksheets:
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        count, header_row, mapping = _find_header(rows)
        if header_row is None or "email" not in mapping:
            continue
        col = mapping["email"]
        usable = sum(1 for r in rows[header_row:] if col < len(r) and emails_in(r[col]))
        if not usable:
            continue
        found.append((count, usable, ws, rows, header_row, mapping))
    if not found:
        wb.close()
        out["problem"] = ("No sheet has a recognisable header row with an Email column. "
                          "Start from the contact template, or name the columns as it does.")
        return out
    # the main sheet first: the one whose header is understood best
    found.sort(key=lambda f: (-f[0], -f[1]))
    seen = {}
    unmatched, have = [], set()
    for count, usable, ws, rows, header_row, mapping in found:
        matrix = _owner_matrix(rows, header_row, mapping)
        out["sheets"].append(ws.title)
        have |= set(mapping)
        header = rows[header_row - 1]
        matched = set(mapping.values()) | set(matrix)
        unmatched += [" ".join(_cell(v).split()) for c, v in enumerate(header)
                      if _cell(v) and c not in matched and " ".join(_cell(v).split()) not in unmatched]
        for contact in _sheet_contacts(ws, rows, header_row, mapping, matrix):
            email = contact["email"]
            where = "row %d" % contact["row"] if contact["sheet"] == out["sheets"][0] else \
                "row %d of %s" % (contact["row"], contact["sheet"])
            if not email:
                contact["problems"].append("no email address")
            elif not EMAIL_RE.match(email):
                contact["problems"].append("email address looks wrong")
            elif email in seen:
                contact["problems"].append("same email as %s" % seen[email])
            else:
                seen[email] = where
            if not contact["last"] and not contact["first"]:
                contact["problems"].append("no name")
            out["contacts"].append(contact)
    wb.close()
    out["sheet"] = out["sheets"][0]
    out["unmatched_headers"] = unmatched
    out["missing_columns"] = [f for f in REQUIRED if f not in have and not (f == "name" and have & {"last", "person"})]
    out["worked_out"] = [f for f in WORKED_OUT if f not in have]
    speaking = set(directory().get("german_speaking", ["AT", "DE", "CH", "LI"]))
    # A row that leaves a column empty is noted on the row; a list without the
    # column at all is noted once, on the Setup tab (worked_out).
    for contact in out["contacts"]:
        if not contact["language"]:
            if contact["country"] and contact["country"] != "MULTI" and contact["country"] not in speaking:
                contact["language"] = "en"
                if "language" in have:
                    contact["assumed"].append("language not given, English assumed from the country")
            else:
                contact["language"] = "de"
                if "language" in have:
                    contact["assumed"].append("language not given, German assumed")
        if not contact["form"]:
            contact["form"] = "sie"
            if "form" in have:
                contact["assumed"].append("Sie/du not given, Sie assumed")
        if contact["owner"]:
            for code in re.split(r"[,;/&+ ]+", contact["owner"]):
                if code:
                    out["owners"][code] = out["owners"].get(code, 0) + 1
    out["ok"] = True
    return out


# ---- greeting -----------------------------------------------------------------

def greeting(c):
    """The salutation line the list gives by rule. Claude refines it when it
    adjusts the email to the person (a first name usually says Herr or Frau),
    but this is what stands when the list is clear or Claude is not asked:
    a wrong 'Herr' or a missing 'Dr.' is the mistake a recipient does not
    forgive. Returns (line, warning or '')."""
    if c.get("greeting"):
        line = c["greeting"].strip()
        return (line if line.endswith(",") else line + ","), ""
    first, last, title, g = c.get("first", ""), c.get("last", ""), c.get("title", ""), c.get("gender", "")
    titled_last = " ".join(p for p in (title, last) if p)
    if c.get("language") == "en":
        if not first and not last:
            return "Dear Sir or Madam,", "no contact person"
        if c.get("form") == "du":
            return ("Dear %s," % first, "") if first else ("Dear %s," % titled_last, "no first name")
        if title.lower().rstrip(".") in ("dr", "prof", "prof dr"):
            return "Dear %s %s," % (title.rstrip(".") if title.lower().startswith("dr") else title, last), ""
        if g and last:
            return "Dear %s %s," % ("Mr" if g == "m" else "Ms", last), ""
        return "Dear %s," % " ".join(p for p in (first, last) if p), "Mr/Ms not given"
    if not first and not last:
        return "Sehr geehrte Damen und Herren,", "no contact person"
    if c.get("form") == "du":
        if not first:
            return "Guten Tag %s," % titled_last, "no first name"
        if g:
            return "%s %s," % ("Lieber" if g == "m" else "Liebe", first), ""
        return "Hallo %s," % first, ""
    if g and last:
        return "%s %s," % ("Sehr geehrter Herr" if g == "m" else "Sehr geehrte Frau", titled_last), ""
    return ("Guten Tag %s," % " ".join(p for p in (title, first, last) if p)), "Herr/Frau not given"


# ---- whose contact this is ----------------------------------------------------------

def initials(name):
    """'Vsevolod Kornilov' -> 'VK'; 'Kornilov, Vsevolod' -> 'VK'."""
    name = str(name or "").strip()
    if "," in name:
        last, _, first = name.partition(",")
        name = first + " " + last
    words = [w for w in re.split(r"[\s._-]+", _strip_titles(name)[0]) if w and w[0].isalpha()]
    if len(words) < 2:
        return ""
    return (words[0][0] + words[-1][0]).upper()


def code_for(name, email, known=()):
    """The user's short form in the Zuständigkeit column, from the team list,
    else from their name, else from their address (first.last@)."""
    for person in directory().get("team", []):
        if person.get("name") and (_ascii(person["name"]) == _ascii(name) or
                                   _ascii(" ".join(reversed(person["name"].split()))) == _ascii(name)):
            return person["code"]
    local = (email or "").split("@")[0]
    for candidate in (initials(name), initials(local) if re.search(r"[._-]", local) else ""):
        if candidate and (not known or candidate in known):
            return candidate
    return initials(name) or (initials(local) if re.search(r"[._-]", local) else "")


def _tokens(text):
    return {t for t in _norm(text).split() if len(t) > 1}


def owner_codes(text):
    return [t for t in re.split(r"[,;/&+\s]+", str(text or "").upper()) if CODE_RE.match(t)]


def is_mine(contact, owner_name, my_email, my_code="", unassigned_mine=True):
    owner = contact.get("owner", "")
    if not owner:
        return unassigned_mine
    if my_email and my_email.lower() in owner.lower():
        return True
    code = (my_code or "").strip().upper()
    if code:
        for theirs in owner_codes(owner):
            if theirs == code or (len(theirs) == 3 and len(code) == 2 and theirs[:2] == code) or \
                    (len(code) == 3 and theirs == code[:2]):
                return True
    mine, theirs = _tokens(owner_name), _tokens(owner)
    return bool(mine and theirs and (mine <= theirs or theirs <= mine))


def uses_codes(contacts):
    """Does the list name the person responsible by short form (VK, KL), as
    AGC's own list does, rather than by name?"""
    owners = [c["owner"] for c in contacts if c.get("owner")]
    return bool(owners) and sum(1 for o in owners if owner_codes(o)) * 2 > len(owners)


def unassigned_is_mine(contacts, settings):
    """A contact with nobody responsible: in a list kept by short forms that
    is nobody's, so it is left out unless the user says otherwise; in a list
    of names (the template) it is whoever runs the round."""
    value = settings.get("unassigned_mine")
    return (not uses_codes(contacts)) if value is None else bool(value)


# ---- who is in this round ---------------------------------------------------------

def _outside(c, filters):
    """Which of the user's own filters leaves this contact out, or ''."""
    countries = filters.get("countries") or []
    if countries:
        country = c.get("country") or "none"
        if country == "MULTI":
            country = "multi"
        if country not in countries and not (country == "multi" and set(countries) & {"AT", "DE", "CH"}):
            return "region"
    states = filters.get("states") or []
    if states and c.get("state") not in states:
        return "region"
    if filters.get("priorities") and (c.get("priority") or "none") not in filters["priorities"]:
        return "priority"
    if filters.get("relevance") and (c.get("relevance") or "none") not in filters["relevance"]:
        return "relevance"
    if filters.get("types") and c.get("type") not in filters["types"]:
        return "type"
    return ""


def select(contacts, settings, my_email, history, suppression, overrides=None, today=None):
    """Mark every contact with included (bool) and reason (why, or why not).

    overrides: {contact id: True/False} set by hand on the Contacts screen. A
    hand override can bring in someone who is not due yet, or whose address or
    contact person is missing (Claude then looks it up), but never someone who
    asked not to be contacted."""
    overrides = overrides or {}
    today = today or dt.date.today()
    rhythm = settings.get("schedule", {}).get("mode", "manual")
    default_cadence = rhythm if rhythm in CADENCE_DAYS else ""
    cap = int(settings.get("batch_size") or 50)
    filters = settings.get("filters") or {}
    priorities = directory().get("priorities", {})
    unassigned_mine = unassigned_is_mine(contacts, settings)
    taken = 0
    out = []
    for c in contacts:
        c = dict(c)
        past = history.get(c["email"], {}) if c["email"] else {}
        c["last_sent"] = past.get("last_sent", "")
        c["last_outcome"] = past.get("last_outcome", "")
        c["last_summary"] = past.get("last_summary", "")
        c["research"] = []
        included, reason, hard = True, "due", False
        missing = [RESEARCHABLE[p] for p in c["problems"] if p in RESEARCHABLE]
        hard_problems = [p for p in c["problems"] if p not in RESEARCHABLE]
        if hard_problems:
            included, reason, hard = False, hard_problems[0], True
        elif c["dnc"]:
            included, reason, hard = False, "marked do not contact in the list", True
        elif c["email"] in suppression:
            included, reason, hard = False, "asked not to be contacted", True
        elif missing:
            if not c["company"]:
                included, reason, hard = False, c["problems"][0], True
            else:
                c["research"] = sorted(set(missing))
                included = False
                reason = c["problems"][0] + ": tick to have Claude look it up"
        elif not is_mine(c, settings.get("owner_name", ""), my_email, settings.get("owner_code", ""),
                         unassigned_mine):
            included, reason = False, ("belongs to %s" % c["owner"] if c.get("owner")
                                       else "nobody is responsible for this contact in the list")
        elif c["cadence"] == "never":
            included, reason = False, "cadence set to never"
        if included and _outside(c, filters):
            included, reason = False, "outside your %s filter" % _outside(c, filters)
        if included:
            last = c["last_sent"][:10] or c.get("first_contact_date", "")
            if last:
                days = (today - dt.date.fromisoformat(last)).days
                prio_days = priorities.get(c.get("priority"), {}).get("days") or 0
                need = CADENCE_DAYS.get(c["cadence"] or default_cadence) or prio_days or RECENT_GUARD_DAYS
                if days < need:
                    included, reason = False, ("written to %d days ago" % days if c["last_sent"]
                                               else "first contacted %d days ago according to the list" % days)
            elif not c["last_sent"]:
                reason = "never written to"
        forced = overrides.get(c["id"])
        if not hard and forced is not None and forced != included:
            included = forced
            if forced and c["research"]:
                reason = "Claude will look up the missing details"
            else:
                reason = "added by hand" if forced else "left out by hand"
        if included:
            if taken >= cap:
                included, reason = False, "over the batch size of %d, next round" % cap
            else:
                taken += 1
        c["included"], c["reason"], c["locked"] = included, reason, hard
        out.append(c)
    return out
